"""Build the editable plan and technical figures from versioned project sources."""
from pathlib import Path
import importlib.util
import json
import re
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, FancyArrowPatch
from docx.shared import Cm

ROOT=Path(__file__).resolve().parent
PROJECT=ROOT.parents[1]
FIG=ROOT/'figures'
FIG.mkdir(exist_ok=True)
spec=importlib.util.spec_from_file_location('document_style',PROJECT/'docs/implementation/build_documents.py')
style=importlib.util.module_from_spec(spec)
spec.loader.exec_module(style)
plt.rcParams.update({'font.family':['Microsoft YaHei','DejaVu Sans'],'font.size':11,'axes.unicode_minus':False,'figure.facecolor':'white'})

def canvas(height=3.2):
    fig,ax=plt.subplots(figsize=(9,height));ax.set_xlim(0,10);ax.set_ylim(0,4);ax.axis('off');return fig,ax
def save(fig,name):
    fig.savefig(FIG/(name+'.png'),dpi=190,bbox_inches='tight',pad_inches=.15);plt.close(fig)
def arrow(ax,a,b,color='#385a73'):
    ax.add_patch(FancyArrowPatch(a,b,arrowstyle='-|>',mutation_scale=16,color=color,lw=1.5))
def box(ax,x,y,w,h,text):
    ax.add_patch(Rectangle((x,y),w,h,facecolor='#edf3f7',edgecolor='#758b9a',lw=1));ax.text(x+w/2,y+h/2,text,ha='center',va='center')

fig,ax=canvas(3.0)
box(ax,.2,2.65,2.6,1,'航班与跑道天气\n已确认输入快照')
box(ax,3.65,2.65,2.7,1,'计算内核\n空气状态  受力  轨迹')
box(ax,7.2,2.65,2.6,1,'工程约束求解\nV1  FLEX  TOGA限重')
box(ax,.2,.4,2.6,1,'版本化参数包\n单位  范围  假设')
box(ax,3.65,.4,2.7,1,'唯一结果对象\n状态  余量  诊断')
box(ax,7.2,.4,2.6,1,'页面 图形 报告\n输入变化即失效')
arrow(ax,(2.8,3.15),(3.65,3.15));arrow(ax,(6.35,3.15),(7.2,3.15));arrow(ax,(8.5,2.65),(5.3,1.4));arrow(ax,(2.8,.9),(3.65,2.8));arrow(ax,(6.35,.9),(7.2,.9))
save(fig,'architecture')

fig,ax=canvas(2.9)
ax.plot([.4,9.7],[.65,1.3],color='#465761',lw=3)
box(ax,4.3,1.35,1.5,.55,'质点 m')
arrow(ax,(5.05,1.65),(7.55,1.82),'#157d73');ax.text(7.7,1.9,'推力 T',color='#157d73')
arrow(ax,(4.8,1.7),(2.4,1.53),'#a16023');ax.text(.35,1.95,'阻力 D + 摩擦 μ(W−L)',color='#a16023')
arrow(ax,(5.05,1.7),(5.05,2.9));ax.text(5.22,2.75,'升力 L')
arrow(ax,(5.05,1.7),(5.05,.18));ax.text(5.22,.1,'重力 W')
ax.text(.55,3.8,'地面运动  m·dv/dt = T − D − μ·max(W−L,0) − W·slope/100')
ax.text(.55,3.3,'空气动力使用 TAS；跑道位移使用地速；图为受力示意')
save(fig,'forces')

fig,ax=canvas(2.7)
ax.add_patch(Rectangle((1,2.65),6,.5,facecolor='#e7ebee',edgecolor='#66737d'))
ax.add_patch(Rectangle((7,2.65),1,.5,facecolor='#f4dfba',edgecolor='#ad7a30'))
ax.add_patch(Rectangle((7,3.2),2,.45,facecolor='#d7eee8',edgecolor='#337b69'))
ax.text(3.8,2.9,'起飞跑道',ha='center');ax.text(7.5,2.9,'停止道',ha='center',fontsize=9);ax.text(8,3.43,'净空道',ha='center',fontsize=9)
for y,end,label,color in [(2.05,7,'TORA 约束起飞地面距离','#436e93'),(1.35,9,'TODA 约束继续起飞距离','#287866'),(.65,8,'ASDA 约束加速停止距离','#ac7b32')]:
    ax.annotate('',xy=(end,y),xytext=(2,y),arrowprops={'arrowstyle':'<->','color':color});ax.text(2.1,y+.1,label,color=color)
ax.plot([2,2],[.3,3.1],ls='--',color='#ac4844');ax.text(2.05,3.6,'入口扣减后，从同一新起点比较三项距离',fontsize=10)
save(fig,'runway')

sample_path=ROOT/'evidence/model-sample.json'
if sample_path.exists():
    result=json.loads(sample_path.read_text('utf8'));s=result['solution']
    fig,ax=plt.subplots(figsize=(9,3.4))
    for key,label,color in [('accelerateStop','加速停止','#ad7932'),('continueTakeoff','单发继续起飞','#287866')]:
        points=s['trajectory'][key];ax.plot([p['distanceM'] for p in points],[p['speedKt'] for p in points],label=label,color=color,lw=2)
    ax.set_xlabel('地面距离 m');ax.set_ylabel('校准空速 kt');ax.set_title('工程模型轨迹  非机型校准结果',fontsize=12);ax.grid(alpha=.22);ax.legend(frameon=False);fig.tight_layout();save(fig,'trajectory')

figures={
 '三 系统结构':('architecture','图1 起飞工程模型的输入 计算和输出关系'),
 '四 运动模型与速度求解':('forces','图2 跑道纵向受力模型'),
 '五 跑道约束与限制重量':('runway','图3 三类声明距离的不同约束对象 示意非比例'),
 '七 界面图形与报告':('trajectory','图4 当前模型实际求解的两条轨迹')}
d=style.new_doc('A330EFB 起飞性能工程演示实施方案')
d.core_properties.subject='A330EFB 1.2 起飞工程模型设计 实施与验收'
section=''
def insert_figure():
    if section in figures:
        file,caption=figures[section]
        if (FIG/(file+'.png')).exists():
            p=d.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(FIG/(file+'.png')),width=Cm(16.2))
            p=d.add_paragraph(caption,style='Caption');p.paragraph_format.first_line_indent=Cm(0)
for line in (ROOT/'起飞性能实施方案.md').read_text('utf8').splitlines()[1:]:
    if not line.strip():continue
    if line.startswith('## '):
        insert_figure();section=line[3:];d.add_heading(section,1)
    else:style.add_text(d,line)
insert_figure()
d.add_heading('参考资料与配套文件',1)
for text in [
 'Headwind 与 FlyByWire 源码采用项目锁定版本。资料核对见 docs/takeoff/资料与模型边界.md，实际参数与方法见 模型说明.md。',
 'Airbus Control your speed at take off  https://safetyfirst.airbus.com/control-your-speed-at-take-off/',
 'Airbus A330 机场规划资料 2025年12月版  https://www.aircraft.airbus.com/sites/g/files/jlcbta126/files/2025-12/AC_A330_20251201.pdf',
 '任务状态见 docs/takeoff/tasks.json，操作和程序画面见 操作说明.md，完成情况以 验收记录.md 为准。']:
    style.add_text(d,text)
out=PROJECT/'output/word/A330EFB起飞性能工程演示实施方案.docx';d.save(out)
print(out)
