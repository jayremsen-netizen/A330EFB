"""Build the implementation documents from the versioned plan and task register."""
from pathlib import Path
import json, re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parents[1]
OUT = PROJECT / 'output/word'

def font(style, size, chinese='宋体', bold=False):
    style.font.name='Arial'; style.font.size=Pt(size)
    style.font.bold=bold; style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr()
    rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
    for k,v in [('ascii','Arial'),('hAnsi','Arial'),('eastAsia',chinese)]:rf.set(qn('w:'+k),v)

def new_doc(title):
    d=Document(); s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(2.0);s.bottom_margin=Cm(2.0);s.left_margin=Cm(2.3);s.right_margin=Cm(2.1)
    for name,size,cn,b in [('Normal',11,'宋体',False),('Title',23,'黑体',False),('Heading 1',15,'黑体',True),('Heading 2',12,'黑体',True),('Caption',9,'宋体',False),('Footer',9,'宋体',False)]:
        st=d.styles[name];font(st,size,cn,b)
        for x in list(st.element.iter(qn('w:pBdr'))):x.getparent().remove(x)
        pf=st.paragraph_format;pf.space_after=Pt(6);pf.line_spacing=1.25;pf.widow_control=True
        if name.startswith('Heading'):pf.space_before=Pt(12);pf.keep_with_next=True
    d.styles['Normal'].paragraph_format.first_line_indent=Pt(22)
    d.styles['Title'].paragraph_format.space_after=Pt(12)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    d.core_properties.title=title;d.core_properties.author='A330EFB';d.core_properties.last_modified_by='A330EFB'
    d.core_properties.subject='A330EFB 1.1 投标演示改造与验收';d.core_properties.comments=''
    d.add_paragraph(title,style='Title')
    return d

def add_text(d,text):
    p=d.add_paragraph()
    # Keep prose editable, including explicit emphasis from the source.
    for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
        p.add_run(part).bold=(i%2==1)
    return p

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    register=json.loads((ROOT/'tasks.json').read_text('utf8'))
    source=(ROOT/'方案.md').read_text('utf8')
    d=new_doc('A330EFB 投标演示版完善方案')
    for line in source.splitlines()[1:]:
        if not line.strip():continue
        if line.startswith('## '):d.add_heading(line[3:],1)
        else:add_text(d,line)
    d.save(OUT/'A330EFB投标演示版完善方案.docx')

    lines=['# A330EFB 投标演示版实施任务书','版本 1.1　2026年10月4日',
      '本任务书对应完善方案确定的投标演示范围。实施顺序由任务依赖确定，各项完成状态以 tasks.json 为准；验收证据记录具体程序行为、测试结果和交付提交。目标是使主要工作流在断网和现场自由操作条件下保持可解释、可恢复。',
      '## 一 工作分解与责任',
      '实施工作按领域规则、宿主适配、存储、资料、运行模型、演示和交付划分。代码修改由实施人员完成并提交对应检查；复核依据实际界面和测试记录确认结果。任务之间传递源码、接口和已通过的检查，未完成前置条件时不得仅凭页面外观宣布后续任务完成。',
      '## 二 分项实施任务']
    labels={'todo':'待实施','in_progress':'实施中','done':'已完成'}
    for t in register['tasks']:
        lines += ['### '+t['id']+' '+t['title'],t['scope'],
          '前置任务：'+('、'.join(t['depends']) or '无')+'。当前状态：'+labels.get(t['status'],t['status'])+'。',
          '实施位置：'+'；'.join(t['files'])+'。', '完成条件：'+t['acceptance']]
        if t.get('evidence'):lines.append('验收证据：'+t['evidence'])
    lines += ['## 三 里程碑与推进条件',
      'M1 完成 T01，形成实施范围与基线。M2 完成 T02 至 T05，消除负结果、数量不一致及数据丢失风险。M3 完成 T06 至 T09，使资料、地服和故障有可实际操作的本地工作流。M4 完成 T10 至 T12，形成复核证据并交付远端仓库。',
      '里程碑按完成条件推进。本轮持续实施至验收通过，不以未经测量的工时承诺替代完成证据；发生阻碍时，记录受影响任务、已经尝试的处置以及仍需解决的依赖。',
      '## 四 验收场景与观察要求',
      'A01 航班准备：导入默认航班，人数 250 改为 260 时 TOW 增加 1040 kg；恢复输入、确认计划、计算、保存历史并导出实际报告。',
      'A02 计算边界：130 吨着陆不得输出负距离；190 吨基准可计算；不支持范围显示原因并清除旧结果。CONF 3 疑点区间拒绝输出。',
      'A03 质量与单位：50000 kg 额外货物的超限输入不能确认；44836 kg 货舱边界保持一致；kg 与 lb 请求返回对应数值，写入后反向读取保持原质量。',
      'A04 燃油：5000 kg 和 30000 kg 分配满足总量守恒，计划初始化与目标使用相同规则；地面油量变化使旧报告不可导出，采用新计划后可以重算。',
      'A05 草稿恢复：清空人数后刷新，航班号和其他输入仍保留；补齐人数后可以确认。异常历史和结果被隔离，不触发整页错误，恢复备份可导出。',
      'A06 保存冲突：窗口 A 修改人数，窗口 B 修改航班号，旧窗口不得静默覆盖 A；加载最新记录后继续编辑。检查演示前缀不影响普通航班。',
      'A07 资料阅读：阻断外网仍能检索内置图件、缩放、旋转和收藏；导入 PDF 或图片可以阅读，异常文件不破坏已有页面。',
      'A08 天气采用：选择带来源时间的样例，明确采用后更新计划并使旧结果失效；再次确认计算后报告带当前输入。',
      'A09 地服与故障：电源连接、加油和登机经历执行到完成；条件不满足的推出被拒绝；故障影响操作，解除后能够恢复；重复命令和会话重置不产生旧动作残留。',
      'A10 自动演示：四套场景独立运行，支持暂停、接管、重播、定位和下载。全流程在无外网条件下完成，错误有可读反馈与恢复入口。',
      'A11 可复现交付：按 README 从仓库准备依赖和固定上游，Windows 与 Ubuntu 构建、单元、浏览器检查及仓库审查通过。',
      'A12 文档对应：现行说明引用正确版本章节；方案、任务、操作补充和完成记录与程序行为一致；Word 文件全部页面通过排版复核。',
      '## 五 测试证据和完成判定',
      '测试结果使用机器可读 JSON 保存，浏览器场景保留关键画面。证据包括软件版本、测试时刻、命令、输入、预期、实际结果及失败原因。修复导致现有算例改变时，必须说明业务规则变化，不能直接改成新的输出数值作为通过条件。',
      '任务完成需要代码或文档成果、对应验收场景和可定位证据同时具备。性能资料尚未校核的范围以明确限制验收，不能登记为全范围性能算法完成。',
      '## 六 交付目录和维护安排',
      'docs/implementation 保存方案、任务源文件及生成脚本；output/word 保存可编辑文档；tests 保存回归入口；docs/verification 保存最终验收摘要；local-extensions 与 local-efb 保存实现。资料源文件与代码一起版本化，上游保持 Git 引用。',
      '发布后以最终提交和文档版本识别交付内容。后续增加真实宿主或可靠性能资料时，另立接入任务和回归基线，沿用现有单位、数据身份、命令回执及结果失效约定。']
    (ROOT/'实施任务书.md').write_text('\n\n'.join(lines)+'\n','utf8')
    d=new_doc('A330EFB 投标演示版实施任务书')
    for line in lines[1:]:
        if line.startswith('### '):d.add_heading(line[4:],2)
        elif line.startswith('## '):d.add_heading(line[3:],1)
        else:add_text(d,line)
    d.save(OUT/'A330EFB投标演示版实施任务书.docx')
    print('Created implementation plan and task book.')

if __name__=='__main__':main()
