"""Build the editable technical manual with continuous pagination and real Word styles."""
from pathlib import Path
import json,re,argparse,collections
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

ROOT=Path(__file__).resolve().parent;WORK=ROOT.parent
OUTPUT=WORK/'output/word/A330电子飞行包系统技术方案与操作说明书.docx'

def east(style,cn='宋体',latin='Times New Roman',size=11):
    style.font.name=latin;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
    for attr,v in [('ascii',latin),('hAnsi',latin),('eastAsia',cn)]:fonts.set(qn('w:'+attr),v)

def field(p,instruction,result=''):
    a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');p.add_run()._r.append(a)
    a=OxmlElement('w:instrText');a.set(qn('xml:space'),'preserve');a.text=instruction;p.add_run()._r.append(a)
    a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'separate');p.add_run()._r.append(a)
    p.add_run(result)
    a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'end');p.add_run()._r.append(a)

def bookmark(p,name,index):
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(index));a.set(qn('w:name'),name);p._p.insert(0,a)
    a=OxmlElement('w:bookmarkEnd');a.set(qn('w:id'),str(index));p._p.append(a)

def link(p,label,target,internal=False):
    h=OxmlElement('w:hyperlink')
    if internal:h.set(qn('w:anchor'),target)
    else:h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');props=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'000000');props.append(c)
    r.append(props);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def title_text(text):
    return re.sub(r'[、，。；：:（）()《》/·—–,;!?]', ' ',text).replace('  ',' ').strip()

def build(page_map=None):
    book=json.loads((ROOT/'manuscript.json').read_text('utf-8'));doc=Document();page_map=page_map or {}
    sec=doc.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21)
    sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.25);sec.right_margin=Cm(2.05)
    sec.header_distance=Cm(.8);sec.footer_distance=Cm(.9);sec.different_first_page_header_footer=True
    styles=doc.styles
    for sty in styles:
        if sty.type==WD_STYLE_TYPE.PARAGRAPH:
            for border in list(sty.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    east(styles['Normal']);pf=styles['Normal'].paragraph_format
    pf.line_spacing=1.28;pf.space_after=Pt(5);pf.widow_control=True;pf.first_line_indent=Pt(22)
    for name,size,cn in [('Title',25,'黑体'),('Subtitle',13,'宋体'),('Heading 1',17,'黑体'),('Heading 2',13,'黑体'),('Heading 3',11.5,'黑体')]:
        s=styles[name];east(s,cn,'Arial',size);s.font.bold=name not in ('Title','Subtitle');s.font.italic=False;f=s.paragraph_format;f.first_line_indent=Pt(0);f.space_before=Pt(13 if name.startswith('Heading') else 0);f.space_after=Pt(8);f.keep_with_next=True;f.line_spacing=1.1
    styles['Heading 1'].paragraph_format.page_break_before=True
    front=styles.add_style('FrontHeading',WD_STYLE_TYPE.PARAGRAPH);front.base_style=styles['Heading 1']
    outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),'9');front.element.get_or_add_pPr().append(outline)
    for name,size in [('Caption',9),('RefNote',8.5),('Equation',10.5),('Step',10.5),('TOC 1',10.5),('TOC 2',10),('Code',9),('TableText',9)]:
        s=styles[name] if name in styles else styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        east(s,'宋体','Times New Roman',size);f=s.paragraph_format;f.first_line_indent=Pt(0);f.space_before=Pt(0);f.space_after=Pt(4);f.line_spacing=1.18;f.widow_control=True
    styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    styles['Caption'].font.italic=False
    styles['Caption'].font.bold=False
    styles['Equation'].paragraph_format.left_indent=Cm(.35)
    styles['Step'].paragraph_format.left_indent=Cm(.6);styles['Step'].paragraph_format.first_line_indent=Cm(-.6)
    styles['TOC 1'].font.bold=True;styles['TOC 1'].paragraph_format.space_before=Pt(4)
    styles['TOC 1'].paragraph_format.space_after=Pt(2)
    styles['TOC 2'].paragraph_format.left_indent=Cm(.55);styles['TOC 2'].paragraph_format.space_after=Pt(1)
    for name in ['TOC 1','TOC 2']:styles[name].paragraph_format.tab_stops.add_tab_stop(Cm(16.3),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    styles['RefNote'].font.color.rgb=RGBColor.from_string('555555')
    east(styles['Header'],'宋体','Arial',8);east(styles['Footer'],'宋体','Times New Roman',9)
    header=sec.header.paragraphs[0];header.text='A330EFB  技术方案与操作说明书';header.paragraph_format.first_line_indent=Pt(0)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Pt(0)
    field(footer,' PAGE ','1')
    doc.core_properties.title=book['title'];doc.core_properties.subject='A330EFB 系统设计与操作';doc.core_properties.author='A330EFB';doc.core_properties.keywords='A330EFB Headwind flyPad 系统设计 操作说明'
    doc.core_properties.comments='';doc.core_properties.last_modified_by='A330EFB'
    settings=doc.settings.element;upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'false');settings.append(upd)
    p=doc.add_paragraph(style='Title');p.paragraph_format.space_before=Cm(4.2);p.add_run('A330电子飞行包系统\n技术方案与操作说明书')
    p=doc.add_paragraph('A330EFB',style='Subtitle');p.paragraph_format.space_before=Cm(.8)
    p=doc.add_paragraph('系统架构与模块设计\n关键技术与操作使用',style='Subtitle');p.paragraph_format.space_before=Cm(.8)
    p=doc.add_paragraph('文档版本 2.0\n2026年10月',style='Normal');p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Cm(4)
    doc.add_paragraph('适用范围',style='FrontHeading')
    for t in [
      '本说明书面向 A330EFB 的方案评审人员、开发人员及演示操作人员，说明浏览器版本的系统结构、模块职责、接口与数据模型，并给出完整的本地航班操作和自动演示方法。系统以 Headwind A330-941 的 EFB 为基础，复用 FlyByWire 共享组件，在独立工程中增加本地航班管理、重量核对、参考计算及演示控制。',
      '软件当前用于功能演示与开发验证。本地核心流程无需外部账户；模拟器物理状态、在线签派、商业航图等能力取决于对应宿主或资料服务。起飞模块输出固定源码速度表参考值，不提供完整的起飞放行性能结论。各模块的输入条件、已实现行为和扩展接口在相应章节分别说明。',
      '第1至6章说明系统架构、部署、数据和接口；第7至15章说明业务模块；第16至17章说明状态一致性、构建与运行机制；第18至23章给出实际页面操作和典型演示工作流。计算算例使用明确的样例输入，页面截图对应本地运行状态。',
      '本说明书以 A330EFB 软件提交 f528013 和两个固定上游提交为基线：Headwind 41eace79ed442696a6361dc72947954c9a6cf5cb，FlyByWire 1bf4b8edccf84d0fb83d0eb15e42f2c773e09582。在线项目文档用于解释产品概念和页面用途；具体字段、计算方法和本地行为以固定源码及本项目实现为准。',
      '计量单位随字段列明。内部质量使用 kg，速度参考使用 kt，高度使用 ft，跑道声明距离使用 m，气压使用 hPa。界面切换单位不改变内部存储单位。Flight 表示航班输入，Result 表示某次计算快照；计划、地面目标与宿主观测保持独立含义。',
    ]:doc.add_paragraph(t)
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);link(p,'项目源码及构建说明','https://github.com/jayremsen-netizen/A330EFB')
    doc.add_paragraph('目录',style='FrontHeading')
    # A real TOC field, with a usable cached, bookmarked result for Word and preview readers.
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1)
    beg=OxmlElement('w:fldChar');beg.set(qn('w:fldCharType'),'begin');p.add_run()._r.append(beg)
    ins=OxmlElement('w:instrText');ins.set(qn('xml:space'),'preserve');ins.text=' TOC \\o "1-2" \\h \\z \\u ';p.add_run()._r.append(ins)
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');p.add_run()._r.append(sep)
    targets=[]
    for c in book['chapters']:
        label=f"第{c['number']}章 {title_text(c['title'])}";anchor=f"chapter_{c['number']}";targets.append((label,anchor,1))
        for s in c['sections']:
            label=f"第{s['id'].split('.')[1]}节 {title_text(s['title'])}";anchor='section_'+s['id'].replace('.','_');targets.append((label,anchor,2))
    targets.append(('参考资料与源码索引','references',1))
    for label,anchor,level in targets:
        p=doc.add_paragraph(style=f'TOC {level}');link(p,label,anchor,True);p.add_run('\t'+str(page_map.get(anchor,'')))
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');doc.paragraphs[-1].add_run()._r.append(end)
    bm=1;figcounts=collections.Counter();tablecounts=collections.Counter();stats=collections.Counter()
    for c in book['chapters']:
        n=c['number'];p=doc.add_heading(f"第{n}章 {title_text(c['title'])}",1);bookmark(p,f'chapter_{n}',bm);bm+=1
        doc.add_paragraph(c['opening'])
        for s in c['sections']:
            p=doc.add_heading(f"第{s['id'].split('.')[1]}节 {title_text(s['title'])}",2);bookmark(p,'section_'+s['id'].replace('.','_'),bm);bm+=1
            for b in s['blocks']:
                kind=b['type'];stats[kind]+=1
                if kind=='paragraph':doc.add_paragraph(b['text'])
                elif kind=='subheading':
                    p=doc.add_heading(title_text(b['text']),3)
                elif kind=='steps':
                    for i,item in enumerate(b['items'],1):doc.add_paragraph(f'{i}.  {item}',style='Step')
                elif kind=='equations':
                    for label,formula in b['items']:
                        p=doc.add_paragraph(style='Equation');p.add_run(label+'：').bold=True
                        p.add_run(formula)
                elif kind=='code':
                    doc.add_paragraph(b['text'],style='Code')
                elif kind=='figure':
                    path=WORK/b['path'];w,h=Image.open(path).size
                    maxw=16.6;maxh=(8.0 if b.get('compact') else 10.3) if b['kind']=='screen' else 9.6
                    width=min(maxw,maxh*w/h)
                    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True;p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    shape=p.add_run().add_picture(str(path),width=Cm(width));shape._inline.docPr.set('descr',b['caption'])
                    figcounts[n]+=1;doc.add_paragraph(f"图{n}-{figcounts[n]}  {b['caption']}",style='Caption')
                elif kind=='table':
                    tablecounts[n]+=1;p=doc.add_paragraph(f"表{n}-{tablecounts[n]}  {b['caption']}",style='Caption');p.paragraph_format.keep_with_next=True
                    headers=b['headers'];tab=doc.add_table(rows=1,cols=len(headers));tab.autofit=False
                    for i,t in enumerate(headers):tab.rows[0].cells[i].text=str(t)
                    for row in b['rows']:
                        cells=tab.add_row().cells
                        for i,t in enumerate(row):cells[i].text=str(t)
                    pr=tab._tbl.tblPr;borders=OxmlElement('w:tblBorders')
                    for side in ['top','left','bottom','right','insideH','insideV']:
                        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D4D7DA');borders.append(e)
                    pr.append(borders)
                    for ri,row in enumerate(tab.rows):
                        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
                        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
                        for cell in row.cells:
                            tcpr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
                            for side,value in [('top','70'),('bottom','70'),('left','85'),('right','85')]:
                                el=OxmlElement('w:'+side);el.set(qn('w:w'),value);el.set(qn('w:type'),'dxa');m.append(el)
                            tcpr.append(m)
                            for p in cell.paragraphs:
                                p.style=styles['TableText']
                                for r in p.runs:r.bold=(ri==0)
                    doc.add_paragraph().paragraph_format.space_after=Pt(0)
            doc.paragraphs[-1].paragraph_format.keep_with_next=True
            refs=' '.join('['+r+']' for r in s['refs'])
            doc.add_paragraph('参考依据  '+refs,style='RefNote')
    p=doc.add_heading('参考资料与源码索引',1);bookmark(p,'references',bm)
    doc.add_paragraph('下列资料标识与各节末的参考依据对应。GitHub 链接固定到已采用的提交；在线手册内容可能随项目更新。本地路径均相对于 A330EFB 工程根目录，build-common 及参考数据文件由构建脚本生成。')
    for code,(label,source) in book['sources'].items():
        p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.add_run(f'[{code}] {label}').bold=True
        p=doc.add_paragraph(style='RefNote')
        if source.startswith('https://'):link(p,source,source)
        else:p.add_run(source)
    OUTPUT.parent.mkdir(exist_ok=True,parents=True);doc.save(OUTPUT)
    (ROOT/'document-index.json').write_text(json.dumps(dict(targets=targets,figures=sum(figcounts.values()),tables=sum(tablecounts.values()),blocks=stats),ensure_ascii=False,indent=2),'utf-8')
    print(OUTPUT);print(dict(figures=sum(figcounts.values()),tables=sum(tablecounts.values()),sections=sum(len(c['sections']) for c in book['chapters'])))

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--page-map');args=ap.parse_args()
    build(json.loads(Path(args.page_map).read_text('utf-8')) if args.page_map else None)
